"""Reading the client's workbook, writing module outputs, and syncing manual overrides."""
from __future__ import annotations

import logging
import os
import time
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import DB
from .urls import url_key

log = logging.getLogger(__name__)

HEADER_ALIASES = {
    "name": ("institution name",),
    "city": ("city",),
    "state": ("state",),
    "unitid": ("unitid",),
    "website": ("website links",),
    "bacc": ("bacc degree",),
    "emails": ("emails",),
    "notes": ("notes",),
    "carnegie": ("carnegie info",),
    "research": ("research activity designation",),
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
OVERRIDE_FILL = PatternFill("solid", fgColor="DDEBF7")
LEVEL_FILLS = {
    "High": PatternFill("solid", fgColor="C6EFCE"),
    "Medium": PatternFill("solid", fgColor="FFEB9C"),
    "Low": PatternFill("solid", fgColor="FFC7CE"),
}

US_STATES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California", "CO": "Colorado",
    "CT": "Connecticut", "DE": "Delaware", "DC": "District of Columbia", "FL": "Florida", "GA": "Georgia",
    "HI": "Hawaii", "ID": "Idaho", "IL": "Illinois", "IN": "Indiana", "IA": "Iowa", "KS": "Kansas",
    "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine", "MD": "Maryland", "MA": "Massachusetts", "MI": "Michigan",
    "MN": "Minnesota", "MS": "Mississippi", "MO": "Missouri", "MT": "Montana", "NE": "Nebraska", "NV": "Nevada",
    "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico", "NY": "New York", "NC": "North Carolina",
    "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma", "OR": "Oregon", "PA": "Pennsylvania",
    "RI": "Rhode Island", "SC": "South Carolina", "SD": "South Dakota", "TN": "Tennessee", "TX": "Texas",
    "UT": "Utah", "VT": "Vermont", "VA": "Virginia", "WA": "Washington", "WV": "West Virginia", "WI": "Wisconsin",
    "WY": "Wyoming", "PR": "Puerto Rico", "GU": "Guam", "VI": "Virgin Islands", "MH": "Marshall Islands",
    "AS": "American Samoa", "MP": "Northern Mariana Islands", "FM": "Micronesia", "PW": "Palau",
}


@dataclass
class Institution:
    unitid: str
    name: str
    city: str
    state: str
    research: str
    row: int


def as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm_unitid(value) -> str:
    return as_text(value)


def carnegie_label(designation) -> str:
    return "R1" if as_text(designation).lower().startswith("research 1") else "Non-R1"


def _header_map(values) -> dict[str, int]:
    """Header key -> 0-based column index."""
    cols = {}
    for idx, value in enumerate(values):
        h = as_text(value).lower()
        for key, aliases in HEADER_ALIASES.items():
            if key not in cols and any(h.startswith(a) for a in aliases):
                cols[key] = idx
    return cols


def load_institutions(path: Path, sheet: str | None = None) -> list[Institution]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        cols = _header_map(next(rows, ()))
        missing = [k for k in ("unitid", "name") if k not in cols]
        if missing:
            raise ValueError(f"Input sheet is missing columns: {missing}")

        def cell(r, key):
            idx = cols.get(key)
            return as_text(r[idx]) if idx is not None and idx < len(r) else ""

        out = []
        for i, r in enumerate(rows, start=2):
            uid, name = cell(r, "unitid"), cell(r, "name")
            if uid and name:
                out.append(Institution(uid, name, cell(r, "city"), cell(r, "state"), cell(r, "research"), i))
        return out
    finally:
        wb.close()


def _state_name(value: str) -> str:
    v = value.strip()
    return US_STATES.get(v.upper(), v).lower() if len(v) == 2 else v.lower()


def filter_institutions(insts: list[Institution], states=None, unitids=None) -> list[Institution]:
    if states:
        wanted = {_state_name(s) for s in states}
        insts = [i for i in insts if i.state.strip().lower() in wanted]
    if unitids:
        wanted_ids = {str(u).strip() for u in unitids}
        insts = [i for i in insts if i.unitid in wanted_ids]
    return insts


# --- generic tables ------------------------------------------------------------
def _clean(value):
    if isinstance(value, (list, set, tuple)):
        value = "; ".join(str(v) for v in value)
    if isinstance(value, bool):
        return "Y" if value else ""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        return value[:32000]
    return value


def new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def write_sheet(wb: Workbook, title: str, columns: list[str], rows: list[dict],
                level_col: str | None = None, override_col: str | None = None):
    ws = wb.create_sheet(title[:31])
    ws.append(columns)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for r in rows:
        values = []
        for col in columns:
            v = _clean(r.get(col, ""))
            if col == "unitid" and isinstance(v, str) and v.isdigit():
                v = int(v)
            values.append(v)
        ws.append(values)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    if level_col in columns:
        idx = columns.index(level_col) + 1
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            fill = LEVEL_FILLS.get(str(cell.value))
            if fill:
                cell.fill = fill
    if override_col in columns:
        idx = columns.index(override_col) + 1
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell.fill = OVERRIDE_FILL
    for i, col in enumerate(columns, start=1):
        longest = max([len(col)] + [len(str(r.get(col, "") or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
    return ws


def save_workbook(wb: Workbook, path: Path, fallback: bool = True) -> Path | None:
    """Save; if the file is open in Excel, save a timestamped copy instead (fallback=False: return None)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f"~tmp_{path.name}")
    try:
        # write a temp file, then swap it in: stopping a run mid-save never leaves a half-written workbook
        wb.save(tmp)
        os.replace(tmp, path)
        return path
    except PermissionError:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        if not fallback:
            return None
        alt = path.with_name(f"{path.stem}_{time.strftime('%Y%m%d_%H%M%S')}{path.suffix}")
        wb.save(alt)
        log.warning("%s is open in another program - saved to %s instead. Close Excel and re-run with "
                    "--export-only to update the main file.", path.name, alt.name)
        return alt


def read_table(path: Path, sheet: str | None = None) -> list[dict]:
    """Rows of a sheet as dicts (header row = keys). Missing named sheet -> []."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in wb.sheetnames:
                return []
            ws = wb[sheet]
        else:
            ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [as_text(h) for h in next(rows, ())]
        out = []
        for r in rows:
            d = {h: as_text(v) for h, v in zip(header, r) if h}
            if any(d.values()):
                out.append(d)
        return out
    finally:
        wb.close()


def sync_overrides(db: DB, module: str, path: Path, column: str, sheets: tuple[str, ...],
                   edited_column: str = "", baseline: dict | None = None) -> None:
    """Copy the manual override column from an output workbook into the database (later sheets win).

    With edited_column + baseline (the values the module wrote), a value typed over the result column itself
    (e.g. official_url) also counts as an override when the override column is empty.
    """
    if not path.exists():
        return
    try:
        values: dict[str, str] = {}
        for sheet in sheets:
            for r in read_table(path, sheet):
                uid, v = r.get("unitid", ""), r.get(column, "").strip()
                if not v and edited_column and baseline is not None:
                    e = r.get(edited_column, "").strip()
                    if e and url_key(e) != url_key(baseline.get(uid, "") or ""):
                        v = e
                if uid and (v or uid not in values):
                    values[uid] = v
        for uid, v in values.items():
            db.set_override(module, uid, v)
        n = sum(1 for v in values.values() if v)
        if n:
            log.info("Using %d manual override(s) from %s", n, path.name)
    except PermissionError:
        log.warning("Could not read %s (open in Excel?) - manual overrides not refreshed", path.name)


# --- final deliverable ---------------------------------------------------------
def write_final(input_path: Path, sheet: str | None, output_path: Path, fills: dict[str, dict],
                extra_sheets: list[tuple], fallback: bool = True) -> Path | None:
    """Copy the client's workbook, fill columns D/E/F/I/NOTES by unitid, and add extra sheets."""
    wb = load_workbook(input_path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    header = [c.value for c in ws[1]]
    cols = {k: i + 1 for k, i in _header_map(header).items()}
    for key in ("website", "bacc", "emails", "notes", "carnegie"):
        if key not in cols:
            log.warning("Column for '%s' not found in the input sheet - it will not be filled", key)

    for r in range(2, ws.max_row + 1):
        uid = norm_unitid(ws.cell(r, cols["unitid"]).value)
        if not uid:
            continue
        if "carnegie" in cols and "research" in cols:
            ws.cell(r, cols["carnegie"]).value = carnegie_label(ws.cell(r, cols["research"]).value)
        fill = fills.get(uid)
        if not fill:
            continue
        for key in ("website", "bacc", "emails"):
            if key in cols and key in fill:
                ws.cell(r, cols[key]).value = _clean(fill[key]) or None
        if "notes" in cols and fill.get("notes"):
            current = as_text(ws.cell(r, cols["notes"]).value)
            note = _clean(fill["notes"])
            ws.cell(r, cols["notes"]).value = f"{current}; {note}" if current and note not in current else (current or note)

    for title, columns, rows, level_col in extra_sheets:
        if title in wb.sheetnames:
            del wb[title]
        write_sheet(wb, title, columns, rows, level_col=level_col)
    return save_workbook(wb, output_path, fallback=fallback)
